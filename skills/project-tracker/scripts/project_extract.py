# -*- coding: utf-8 -*-
"""project_extract: 마일스톤 추출을 위한 '최소 입력'을 만들어 stdout에 출력.
실제 마일스톤 작성은 이 출력을 읽은 LLM(Jarvis)이 수행하고
project_tracker.py --apply 로 기록한다.

기본(토큰 0 수준): 프로젝트 유형 + 디렉토리 구조 + 문서 파일명 목록.
--excerpt: 우선순위 문서 1~2개의 앞부분 발췌까지 추가(선별 읽기).

사용: python project_extract.py "<폴더명|프로젝트명>" [--excerpt]
"""
import os, sys, argparse, re

PROJECTS_ROOT = r"C:\project"
MD_NAME = "PROJECT.md"
EXCLUDE_DIRS = {".git", "node_modules", ".gradle", ".claude", ".pytest_cache",
                "dist", "build", ".next", ".expo", ".vercel", "__pycache__",
                ".idea", ".vscode", "venv", ".venv"}
DOC_EXT = {".md", ".txt", ".docx", ".xlsx", ".pdf", ".pptx"}
EXCERPT_CHARS = 2000


def resolve(query):
    for d in sorted(os.listdir(PROJECTS_ROOT)):
        full = os.path.join(PROJECTS_ROOT, d)
        if not os.path.isdir(full):
            continue
        if d == query or query in d:
            return d, full
    # PROJECT.md name 매칭
    for d in sorted(os.listdir(PROJECTS_ROOT)):
        md = os.path.join(PROJECTS_ROOT, d, MD_NAME)
        if os.path.isfile(md):
            with open(md, encoding="utf-8") as f:
                head = f.read(500)
            m = re.search(r"^name:\s*(.+)$", head, re.MULTILINE)
            if m and (m.group(1).strip() == query or query in m.group(1)):
                return d, os.path.join(PROJECTS_ROOT, d)
    return None, None


def walk(root, max_depth=3):
    docs, tree = [], []
    is_code = False
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames if d not in EXCLUDE_DIRS]
        depth = dirpath[len(root):].count(os.sep)
        if depth >= max_depth:
            dirnames[:] = []
        rel = os.path.relpath(dirpath, root)
        if rel != ".":
            tree.append("  " * depth + os.path.basename(dirpath) + "/")
        for fn in filenames:
            if fn in ("package.json", "src") or fn.endswith((".ts", ".tsx", ".py")):
                is_code = True
            ext = os.path.splitext(fn)[1].lower()
            if ext in DOC_EXT:
                docs.append(os.path.join(dirpath, fn))
    if os.path.isdir(os.path.join(root, "src")) or \
       os.path.isfile(os.path.join(root, "package.json")):
        is_code = True
    return docs, tree, is_code


def pick_priority(docs, root):
    def score(p):
        base = os.path.basename(p).lower()
        if base == "readme.md":
            return 0
        if base == "claude.md":
            return 1
        if re.search(r"기획|계획|plan|proposal|정의서|요건|기능", base):
            return 2
        return 3
    ranked = sorted(docs, key=lambda p: (score(p), -os.path.getmtime(p)))
    return ranked[:2]


def excerpt(path):
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext in (".md", ".txt"):
            with open(path, encoding="utf-8", errors="replace") as f:
                return f.read(EXCERPT_CHARS)
        if ext == ".docx":
            import docx
            d = docx.Document(path)
            return "\n".join(p.text for p in d.paragraphs)[:EXCERPT_CHARS]
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:2]:
                out.append(f"[시트: {ws.title}]")
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= 15:
                        break
                    out.append("\t".join(str(c) for c in row if c is not None))
            return "\n".join(out)[:EXCERPT_CHARS]
        if ext == ".pdf":
            try:
                import pdfplumber
                with pdfplumber.open(path) as pdf:
                    return "\n".join((pg.extract_text() or "")
                                     for pg in pdf.pages[:3])[:EXCERPT_CHARS]
            except ImportError:
                from PyPDF2 import PdfReader
                r = PdfReader(path)
                return "\n".join((pg.extract_text() or "")
                                 for pg in r.pages[:3])[:EXCERPT_CHARS]
    except Exception as e:
        return f"(추출 실패: {e})"
    return "(지원하지 않는 형식 — 파일명만 활용)"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("query")
    ap.add_argument("--excerpt", action="store_true")
    args = ap.parse_args()

    folder, root = resolve(args.query)
    if not root:
        print(f"프로젝트를 찾을 수 없습니다: {args.query}", file=sys.stderr)
        sys.exit(1)

    docs, tree, is_code = walk(root)
    print(f"# 프로젝트: {folder}")
    print(f"# 유형: {'code(코드베이스)' if is_code else 'docs(문서중심)'}")
    print(f"# 경로: {root}\n")
    print("## 폴더 구조")
    print("\n".join(tree[:60]) if tree else "(하위 폴더 없음)")
    print(f"\n## 문서 파일 ({len(docs)}개)")
    for d in docs[:60]:
        print("  " + os.path.relpath(d, root))

    if args.excerpt and docs:
        print("\n## 핵심 문서 발췌")
        for p in pick_priority(docs, root):
            print(f"\n### {os.path.relpath(p, root)}")
            print(excerpt(p))

    print("\n# 위 정보로 마일스톤(4~8개)을 JSON 배열로 만들어 "
          "project_tracker.py --apply 로 등록하세요.")


if __name__ == "__main__":
    main()
